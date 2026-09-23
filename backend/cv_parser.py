"""
CV Parser Module
Extracts clean, readable text from uploaded CV documents (.pdf, .docx, .txt).
"""

import io
from typing import BinaryIO
import pypdf
import docx


def extract_text_from_pdf(file_stream: BinaryIO) -> str:
    """Extracts text from a PDF file stream using pypdf."""
    reader = pypdf.PdfReader(file_stream)
    text_parts = []
    for page_idx, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text.strip())
    return "\n\n".join(text_parts)


def extract_text_from_docx(file_stream: BinaryIO) -> str:
    """Extracts text from a Word (.docx) file stream."""
    doc = docx.Document(file_stream)
    text_parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())
    # Also extract text from tables if present
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))
    return "\n".join(text_parts)


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Decodes plain text files with fallback encodings."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def parse_cv_document(filename: str, content: bytes) -> str:
    """
    Main parser entry point. Detects file format from filename extension,
    extracts the text, and normalizes it.
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    stream = io.BytesIO(content)

    if ext == "pdf":
        raw_text = extract_text_from_pdf(stream)
    elif ext in ("docx", "doc"):
        try:
            raw_text = extract_text_from_docx(stream)
        except Exception:
            # Fallback if docx structure fails
            raw_text = extract_text_from_txt(content)
    elif ext in ("txt", "text", "md"):
        raw_text = extract_text_from_txt(content)
    else:
        # Generic fallback
        try:
            raw_text = extract_text_from_pdf(stream)
        except Exception:
            raw_text = extract_text_from_txt(content)

    # Normalize whitespace while preserving line structure
    cleaned_lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    normalized_text = "\n".join(cleaned_lines)

    if not normalized_text:
        raise ValueError("Could not extract any readable text from the uploaded CV file.")

    return normalized_text
